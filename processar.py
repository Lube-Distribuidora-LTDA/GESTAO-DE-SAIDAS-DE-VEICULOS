# -*- coding: utf-8 -*-
"""
Gestão de Saída de Veículos — Lube Distribuidora / Araujo
=========================================================
Lê os relatórios "Deslocamentos e Paradas" (Getrak) da pasta
PLANILHAS REFERENCIAS, consolida todas as placas e calcula, por DIA + PLACA:

  coluna 1  data
  coluna 2  placa do veículo
  coluna 3  primeiro horário de cada dia (1ª saída do veículo)
  coluna 4  horário de entrada do motorista (06:30)
  coluna 5  horas de atraso

O indicador é SÓ o atraso:

  horas de atraso = 1ª saída − entrada, quando o veículo sai DEPOIS da entrada
                    zero, quando sai às 06:30 ou antes

SEM TOLERÂNCIA: 06:30 é 06:30. A comparação é exata, ao segundo — uma saída
às 06:30:20 é atraso de 20 segundos.

Saídas: dados.js (lido pelo dashboard), dados.json e Gestao_Saida_Veiculos.xlsx
"""
import json
import unicodedata
from datetime import datetime
from pathlib import Path

import openpyxl

try:                                     # a logo no Excel é só um extra —
    from openpyxl.drawing.image import Image as ImagemExcel
    TEM_IMAGEM = True
except ImportError:                      # se a máquina não tiver Pillow, a
    TEM_IMAGEM = False                   # planilha sai sem logo, sem quebrar

BASE = Path(__file__).resolve().parent


# a pasta das planilhas pode estar dentro do projeto ou um nível acima
def _achar_pasta():
    for candidata in (BASE / "PLANILHAS REFERENCIAS",
                      BASE.parent / "PLANILHAS REFERENCIAS"):
        if candidata.is_dir():
            return candidata
    return BASE / "PLANILHAS REFERENCIAS"


PASTA_PLANILHAS = _achar_pasta()

ENTRADA_PADRAO = "06:30"        # horário de entrada do motorista (toda a frota)
CORTE_NOITE = "18:00:00"        # dia com 1ª saída neste horário ou depois não
                                 # conta no indicador (viagem avulsa, não é a
                                 # rotina da manhã) — "" desativa a regra

DIAS_SEMANA = ["Segunda", "Terça", "Quarta", "Quinta", "Sexta", "Sábado", "Domingo"]
MESES = ["jan", "fev", "mar", "abr", "mai", "jun",
         "jul", "ago", "set", "out", "nov", "dez"]
LOGO = BASE / "IMAGENS" / "logo-lube.png"

# ------------------------------------------------------------- cores da Lube
COR_MARCA = "0E176C"           # azul da marca
COR_MARCA_2 = "16217F"
COR_MARCA_ESCURO = "080C3C"
COR_VERMELHO = "F90A19"        # vermelho da marca
COR_ATRASO = "C52C27"
COR_ATRASO_CLARO = "FBE9E8"
COR_OK = "0F7A35"
COR_OK_CLARO = "E7F3EA"
COR_NEUTRO_CLARO = "EEF1F9"
COR_FAIXA = "F5F7FC"           # zebra das linhas pares
COR_BORDA = "D9DEEC"
COR_TEXTO = "1B1F2E"
COR_TEXTO_SUAVE = "6B7280"
BRANCO = "FFFFFF"


# ---------------------------------------------------------------- utilidades
def sem_acento(txt):
    return "".join(c for c in unicodedata.normalize("NFKD", str(txt))
                   if not unicodedata.combining(c)).lower().strip()


def para_datetime(valor):
    """Aceita datetime do Excel ou texto 'dd/mm/aaaa HH:MM:SS'."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        return valor
    txt = str(valor).strip()
    for formato in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M",
                    "%Y-%m-%d %H:%M:%S", "%d/%m/%Y"):
        try:
            return datetime.strptime(txt, formato)
        except ValueError:
            continue
    return None


def para_numero(valor):
    if valor is None or valor == "":
        return 0.0
    if isinstance(valor, (int, float)):
        return float(valor)
    txt = str(valor).strip().replace(".", "").replace(",", ".")
    try:
        return float(txt)
    except ValueError:
        return 0.0


def dur_para_segundos(valor):
    """'02d 04:29:22' ou '01:21:42' -> segundos."""
    if valor is None or valor == "":
        return 0
    txt = str(valor).strip()
    dias = 0
    if "d" in txt:
        parte_dia, _, txt = txt.partition("d")
        try:
            dias = int(parte_dia.strip())
        except ValueError:
            dias = 0
        txt = txt.strip()
    partes = [p for p in txt.split(":") if p != ""]
    if not partes:
        return dias * 86400
    try:
        nums = [int(p) for p in partes]
    except ValueError:
        return dias * 86400
    while len(nums) < 3:
        nums.append(0)
    h, m, s = nums[:3]
    return dias * 86400 + h * 3600 + m * 60 + s


def hhmmss(segundos):
    """Segundos -> 'HH:MM:SS'."""
    segundos = int(round(abs(segundos)))
    return "%02d:%02d:%02d" % (segundos // 3600,
                               (segundos % 3600) // 60, segundos % 60)


def entrada_em_segundos(hora_txt):
    partes = str(hora_txt).split(":")
    while len(partes) < 3:
        partes.append("0")
    return int(partes[0]) * 3600 + int(partes[1]) * 60 + int(partes[2])


# ------------------------------------------------------------------- leitura
COLUNAS = {
    "evento": ["evento"],
    "cliente": ["cliente"],
    "placa": ["placa", "placa / apelido"],
    "motorista": ["motorista"],
    "inicio": ["hora inicio"],
    "final": ["hora final"],
    "tempo_movimento": ["tempo em movimento"],
    "tempo_ocioso": ["tempo ocioso"],
    "duracao": ["duracao"],
    "hod_inicial": ["hodometro inicial (km)", "hodometro inicial"],
    "hod_final": ["hodometro final (km)", "hodometro final"],
    "distancia": ["distancia (km)", "distancia"],
    "local_inicial": ["local inicial"],
    "local_final": ["local final"],
}


def mapear_cabecalho(linha):
    mapa = {}
    for idx, celula in enumerate(linha):
        if celula is None:
            continue
        nome = sem_acento(celula)
        for chave, apelidos in COLUNAS.items():
            if chave in mapa:
                continue
            if nome in apelidos:
                mapa[chave] = idx
    return mapa


def ler_arquivo(caminho):
    """Devolve a lista de deslocamentos de um relatório Getrak."""
    wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    registros = []
    for ws in wb.worksheets:
        linhas = list(ws.iter_rows(values_only=True))
        mapa, primeira_linha = None, 0
        for i, linha in enumerate(linhas):
            candidato = mapear_cabecalho(linha)
            if "inicio" in candidato and "placa" in candidato:
                mapa, primeira_linha = candidato, i + 1
                break
        if not mapa:
            continue

        # placa do cabeçalho do relatório (usada quando a coluna vem vazia)
        placa_meta = ""
        for linha in linhas[:primeira_linha]:
            for celula in linha:
                if celula and sem_acento(celula).startswith("placa"):
                    placa_meta = str(celula).split(":", 1)[-1].strip()
                    break
            if placa_meta:
                break

        def valor(linha, chave):
            idx = mapa.get(chave)
            return linha[idx] if idx is not None and idx < len(linha) else None

        for linha in linhas[primeira_linha:]:
            if linha is None or all(c is None or c == "" for c in linha):
                continue
            inicio = para_datetime(valor(linha, "inicio"))
            if inicio is None:
                continue
            evento = str(valor(linha, "evento") or "D").strip().upper()
            if evento and evento[0] not in ("D", "P"):
                continue
            placa = str(valor(linha, "placa") or placa_meta or "").strip().upper()
            if not placa:
                continue
            motorista = str(valor(linha, "motorista") or "").strip()
            if sem_acento(motorista) in ("sem motorista", "none", ""):
                motorista = ""       # normalizado — mesmo critério do leitor no navegador
            final = para_datetime(valor(linha, "final"))
            registros.append({
                "arquivo": caminho.name,
                "evento": evento[0] if evento else "D",
                "placa": placa,
                "motorista": motorista,
                "inicio": inicio.strftime("%d/%m/%Y %H:%M:%S"),
                "final": final.strftime("%d/%m/%Y %H:%M:%S") if final else "",
                "data": inicio.date().isoformat(),
                "inicio_seg": inicio.hour * 3600 + inicio.minute * 60 + inicio.second,
                "seg_movimento": dur_para_segundos(valor(linha, "tempo_movimento")),
                "seg_ocioso": dur_para_segundos(valor(linha, "tempo_ocioso")),
                "seg_duracao": dur_para_segundos(valor(linha, "duracao")),
                "km": para_numero(valor(linha, "distancia")),
                "hod_inicial": para_numero(valor(linha, "hod_inicial")),
                "hod_final": para_numero(valor(linha, "hod_final")),
                "local_inicial": str(valor(linha, "local_inicial") or "").strip(),
                "local_final": str(valor(linha, "local_final") or "").strip(),
            })
    wb.close()
    return registros


# ------------------------------------------------------------------- cálculo
def montar_dias(registros):
    """Agrupa por (data, placa) e calcula a 1ª saída e as horas de atraso."""
    grupos = {}
    for r in registros:
        if r["evento"] != "D":          # só deslocamentos contam como saída
            continue
        grupos.setdefault((r["data"], r["placa"]), []).append(r)

    seg_corte = entrada_em_segundos(CORTE_NOITE) if CORTE_NOITE else None

    dias = []
    for (data_iso, placa), itens in sorted(grupos.items()):
        itens.sort(key=lambda x: x["inicio_seg"])
        primeiro = itens[0]
        seg_saida = primeiro["inicio_seg"]
        if seg_corte is not None and seg_saida >= seg_corte:
            continue          # saída à noite — não conta no indicador
        dia = datetime.strptime(data_iso, "%Y-%m-%d").date()
        entrada_txt = ENTRADA_PADRAO
        seg_entrada = entrada_em_segundos(entrada_txt)
        atraso = max(seg_saida - seg_entrada, 0)     # só o atraso interessa

        finais = [i["final"] for i in itens if i["final"]]
        motoristas = [i["motorista"] for i in itens if i["motorista"]]
        dias.append({
            "data": data_iso,
            "data_br": dia.strftime("%d/%m/%Y"),
            "dia_semana": DIAS_SEMANA[dia.weekday()],
            "fim_de_semana": dia.weekday() >= 5,
            "mes": "%s/%d" % (MESES[dia.month - 1], dia.year),
            "placa": placa,
            "motorista": motoristas[0] if motoristas else "",   # "" = não identificado
            "primeira_saida": primeiro["inicio"][-8:],
            "primeira_saida_seg": seg_saida,
            "entrada_motorista": entrada_txt,
            "entrada_seg": seg_entrada,
            "atraso_seg": atraso,
            "atraso_hhmmss": hhmmss(atraso),
            "situacao": "Com atraso" if atraso > 0 else "Sem atraso",
            "ultima_chegada": max(finais)[-8:] if finais else "",
            "viagens": len(itens),
            "km": round(sum(i["km"] for i in itens), 1),
            "seg_movimento": sum(i["seg_movimento"] for i in itens),
            "seg_ocioso": sum(i["seg_ocioso"] for i in itens),
            "seg_duracao": sum(i["seg_duracao"] for i in itens),
            "origem": primeiro["local_inicial"],
        })
    return dias


def montar_resumo(dias):
    if not dias:
        return {}
    total = len(dias)
    atrasadas = [d for d in dias if d["atraso_seg"] > 0]
    seg_atraso = sum(d["atraso_seg"] for d in dias)

    por_placa = {}
    for d in dias:
        p = por_placa.setdefault(d["placa"], {
            "placa": d["placa"], "entrada": d["entrada_motorista"], "motorista": "",
            "dias": 0, "atraso_seg": 0, "dias_atrasados": 0,
            "km": 0.0, "viagens": 0, "saida_seg_total": 0,
        })
        if not p["motorista"] and d["motorista"]:
            p["motorista"] = d["motorista"]
        p["dias"] += 1
        p["atraso_seg"] += d["atraso_seg"]
        p["dias_atrasados"] += 1 if d["atraso_seg"] > 0 else 0
        p["km"] += d["km"]
        p["viagens"] += d["viagens"]
        p["saida_seg_total"] += d["primeira_saida_seg"]
    for p in por_placa.values():
        p["km"] = round(p["km"], 1)
        p["atraso_hhmmss"] = hhmmss(p["atraso_seg"])
        p["atraso_medio_hhmmss"] = hhmmss(
            p["atraso_seg"] / p["dias_atrasados"]) if p["dias_atrasados"] else "00:00:00"
        p["saida_media"] = hhmmss(p["saida_seg_total"] / p["dias"])
        p["aderencia"] = round(
            100 * (p["dias"] - p["dias_atrasados"]) / p["dias"], 1)
        p["pct_atraso"] = round(100 - p["aderencia"], 1)

    return {
        "dias_analisados": total,
        "placas": len(por_placa),
        "periodo_inicio": min(dias, key=lambda d: d["data"])["data_br"],
        "periodo_fim": max(dias, key=lambda d: d["data"])["data_br"],
        "entrada_padrao": ENTRADA_PADRAO,
        "corte_noite": CORTE_NOITE,
        "dias_atrasados": len(atrasadas),
        "dias_sem_atraso": total - len(atrasadas),
        "atraso_total_seg": seg_atraso,
        "atraso_total_hhmmss": hhmmss(seg_atraso),
        "atraso_total_horas": round(seg_atraso / 3600, 2),
        "atraso_medio_hhmmss": hhmmss(seg_atraso / len(atrasadas)) if atrasadas else "00:00:00",
        "atraso_maior_hhmmss": hhmmss(max(d["atraso_seg"] for d in dias)),
        "saida_media": hhmmss(sum(d["primeira_saida_seg"] for d in dias) / total),
        "saida_mais_cedo": min(d["primeira_saida"] for d in dias),
        "saida_mais_tarde": max(d["primeira_saida"] for d in dias),
        "km_total": round(sum(d["km"] for d in dias), 1),
        "viagens_total": sum(d["viagens"] for d in dias),
        "movimento_total_hhmmss": hhmmss(sum(d["seg_movimento"] for d in dias)),
        "ocioso_total_hhmmss": hhmmss(sum(d["seg_ocioso"] for d in dias)),
        "aderencia": round(100 * (total - len(atrasadas)) / total, 1),
        "por_placa": sorted(por_placa.values(), key=lambda x: -x["atraso_seg"]),
    }


def montar_por_semana(dias):
    """Agrega por dia da semana (Segunda..Domingo) — revela se algum dia
    específico da semana costuma concentrar mais atraso que os outros."""
    por_dia = {nome: {"dia_semana": nome, "dias": 0, "dias_atrasados": 0,
                      "atraso_seg": 0, "km": 0.0, "viagens": 0}
               for nome in DIAS_SEMANA}
    for d in dias:
        p = por_dia[d["dia_semana"]]
        p["dias"] += 1
        p["atraso_seg"] += d["atraso_seg"]
        p["dias_atrasados"] += 1 if d["atraso_seg"] > 0 else 0
        p["km"] += d["km"]
        p["viagens"] += d["viagens"]
    saida = []
    for nome in DIAS_SEMANA:
        p = por_dia[nome]
        if not p["dias"]:
            continue
        p["km"] = round(p["km"], 1)
        p["atraso_hhmmss"] = hhmmss(p["atraso_seg"])
        p["atraso_medio_hhmmss"] = hhmmss(
            p["atraso_seg"] / p["dias_atrasados"]) if p["dias_atrasados"] else "00:00:00"
        p["pct_atraso"] = round(100 * p["dias_atrasados"] / p["dias"], 1)
        saida.append(p)
    return saida


# -------------------------------------------------------------------- saídas
def gravar_dados(dias, resumo, registros):
    payload = {
        "gerado_em": datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
        "entrada_padrao": ENTRADA_PADRAO,
        "resumo": resumo,
        "dias": dias,
        "deslocamentos": registros,
    }
    (BASE / "dados.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8")
    # dados.js: permite abrir o dashboard com duplo clique (file://),
    # onde o fetch de .json é bloqueado pelo navegador
    (BASE / "dados.js").write_text(
        "window.DADOS = " + json.dumps(payload, ensure_ascii=False) + ";",
        encoding="utf-8")
    return BASE / "dados.js"


def _fracao_hora(segundos):
    """Segundos-do-dia -> fração de dia, pro Excel entender como horário
    (formato 'hh:mm:ss'). Sempre 0–1, dá a volta à meia-noite."""
    return (segundos % 86400) / 86400.0


def _fracao_duracao(segundos):
    """Segundos (sem limite) -> fração de dia, pro Excel entender como
    DURAÇÃO (formato elapsed-time '[h]:mm:ss', que não trava em 24h)."""
    return segundos / 86400.0


def gravar_excel(dias, resumo):
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.formatting.rule import DataBarRule
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.utils import get_column_letter

    destino = BASE / "Gestao_Saida_Veiculos.xlsx"
    wb = openpyxl.Workbook()

    # ---------------------------------------------------------- estilos
    FONTE = "Calibri"
    fundo_marca = PatternFill("solid", fgColor=COR_MARCA)
    fundo_faixa = PatternFill("solid", fgColor=COR_FAIXA)
    fundo_atraso = PatternFill("solid", fgColor=COR_ATRASO_CLARO)
    fundo_ok = PatternFill("solid", fgColor=COR_OK_CLARO)
    fundo_branco = PatternFill("solid", fgColor=BRANCO)
    branco_negrito = Font(name=FONTE, color=BRANCO, bold=True, size=10.5)
    borda_fina = Side(style="thin", color=COR_BORDA)
    borda_cel = Border(left=borda_fina, right=borda_fina, top=borda_fina, bottom=borda_fina)
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
    esquerda = Alignment(horizontal="left", vertical="center", indent=1)

    def cabecalho_tabela(ws, linha, textos, largura_extra=None):
        """Cabeçalho azul Lube, com ou sem congelamento/autofiltro (chamador decide)."""
        for col, texto in enumerate(textos, 1):
            c = ws.cell(row=linha, column=col, value=texto)
            c.fill, c.font, c.alignment = fundo_marca, branco_negrito, centro
            c.border = borda_cel
        ws.row_dimensions[linha].height = 30

    def preparar_tabela(ws, linha_cab, ncols, ultima_linha):
        """Congela o cabeçalho e liga o autofiltro — a aba fica pronta pra
        quem for aplicar os próprios filtros no Excel."""
        ws.freeze_panes = ws.cell(row=linha_cab + 1, column=1).coordinate
        ws.auto_filter.ref = "A%d:%s%d" % (
            linha_cab, get_column_letter(ncols), ultima_linha)

    def larguras(ws, lista):
        for col, w in enumerate(lista, 1):
            ws.column_dimensions[get_column_letter(col)].width = w

    def secao(ws, linha, texto, ncols):
        c = ws.cell(row=linha, column=1, value=texto.upper())
        c.font = Font(name=FONTE, bold=True, size=11.5, color=COR_MARCA)
        c.alignment = Alignment(horizontal="left", vertical="center")
        for col in range(1, ncols + 1):
            ws.cell(row=linha, column=col).border = Border(bottom=Side(style="medium", color=COR_MARCA))
        ws.row_dimensions[linha].height = 22
        return linha + 1

    def tile_kpi(ws, linha, col0, span, cor, rotulo, valor, nota):
        c1 = ws.cell(row=linha, column=col0)
        ws.merge_cells(start_row=linha, start_column=col0, end_row=linha,
                       end_column=col0 + span - 1)
        c1.value = rotulo.upper()
        c1.font = Font(name=FONTE, bold=True, size=9, color=COR_TEXTO_SUAVE)
        c1.alignment = Alignment(horizontal="left", vertical="bottom", indent=1)
        c1.border = Border(left=Side(style="thick", color=cor))
        ws.row_dimensions[linha].height = 16

        linha2 = linha + 1
        c2 = ws.cell(row=linha2, column=col0)
        ws.merge_cells(start_row=linha2, start_column=col0, end_row=linha2,
                       end_column=col0 + span - 1)
        c2.value = CellRichText(
            TextBlock(InlineFont(b=True, sz="19", color=cor, rFont=FONTE), str(valor)),
            TextBlock(InlineFont(sz="8.5", color=COR_TEXTO_SUAVE, rFont=FONTE),
                      "\n" + str(nota)))
        c2.fill = fundo_faixa
        c2.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
        c2.border = Border(left=Side(style="thick", color=cor))
        ws.row_dimensions[linha2].height = 40
        return linha2 + 1

    # ===================================================== Resumo Executivo
    re_ = wb.active
    re_.title = "Resumo Executivo"
    re_.sheet_view.showGridLines = False
    # colunas B e D um pouco mais largas: levam "Motorista" nas duas
    # mini-tabelas (ranking usa B, piores dias usa D) e não podem truncar
    larguras(re_, [13, 20, 13, 20, 13, 13, 13, 13, 13, 13])

    for linha in (1, 2, 3):
        re_.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=10)
        re_.cell(row=linha, column=1).fill = fundo_marca
    re_.row_dimensions[1].height = 34
    re_.row_dimensions[2].height = 19
    re_.row_dimensions[3].height = 19

    c = re_.cell(row=1, column=1, value="GESTÃO DE SAÍDA DE VEÍCULOS")
    c.font = Font(name=FONTE, bold=True, size=17, color=BRANCO)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    c = re_.cell(row=2, column=1, value="Lube Distribuidora  ·  Araujo")
    c.font = Font(name=FONTE, size=11.5, color="D7DCF5")
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    c = re_.cell(row=3, column=1, value="Período: %s a %s   ·   %d veículo(s)   ·   gerado em %s" % (
        resumo["periodo_inicio"], resumo["periodo_fim"], resumo["placas"],
        datetime.now().strftime("%d/%m/%Y %H:%M")))
    c.font = Font(name=FONTE, size=10, color="AAB4E8")
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    if TEM_IMAGEM and LOGO.is_file():
        try:
            img = ImagemExcel(str(LOGO))
            img.width, img.height = 132, 44
            re_.add_image(img, "I1")
        except Exception:
            pass          # logo é só um extra — nunca trava a geração da planilha

    linha = secao(re_, 5, "Indicadores do período", 10)
    linha += 1                                       # linha 6: espaço

    linha_tiles = linha
    tile_kpi(re_, linha_tiles, 1, 2, COR_ATRASO, "Horas de atraso",
             resumo["atraso_total_hhmmss"],
             "%d de %d dias com atraso" % (resumo["dias_atrasados"], resumo["dias_analisados"]))
    tile_kpi(re_, linha_tiles, 3, 2, COR_MARCA, "Aderência ao horário",
             "%s%%" % str(resumo["aderencia"]).replace(".", ","),
             "%d dias sem atraso" % resumo["dias_sem_atraso"])
    tile_kpi(re_, linha_tiles, 5, 2, COR_VERMELHO, "Atraso médio",
             resumo["atraso_medio_hhmmss"], "média dos dias que atrasaram")
    tile_kpi(re_, linha_tiles, 7, 2, COR_ATRASO, "Maior atraso",
             resumo["atraso_maior_hhmmss"], "no período analisado")
    linha = tile_kpi(re_, linha_tiles, 9, 2, COR_MARCA, "Dias analisados",
                     resumo["dias_analisados"],
                     "%d viagens · %s km" % (resumo["viagens_total"],
                                             str(resumo["km_total"]).replace(".", ",")))

    linha += 1
    linha = secao(re_, linha, "Ranking de atraso por veículo", 10)
    cab_rank = ["Placa", "Motorista", "Dias", "Dias c/ atraso", "Horas de atraso",
                "Atraso médio", "Maior atraso", "% atrasado", "Aderência %", "KM"]
    cabecalho_tabela(re_, linha, cab_rank)
    linha_cab_rank = linha
    linha += 1
    linha_dados_ini = linha
    for p in resumo.get("por_placa", []):
        maior_placa = max((d["atraso_seg"] for d in dias if d["placa"] == p["placa"]), default=0)
        vals = [p["placa"], p["motorista"] if p.get("motorista") else "Não identificado",
                p["dias"], p["dias_atrasados"], _fracao_duracao(p["atraso_seg"]),
                _fracao_duracao(p["atraso_seg"] / p["dias_atrasados"]) if p["dias_atrasados"] else 0,
                _fracao_duracao(maior_placa), p["pct_atraso"] / 100.0, p["aderencia"] / 100.0, p["km"]]
        for col, v in enumerate(vals, 1):
            cel = re_.cell(row=linha, column=col, value=v)
            cel.border = borda_cel
            cel.alignment = Alignment(horizontal="center" if col not in (1, 2) else "left",
                                      vertical="center", indent=1 if col == 2 else 0)
            if col in (5, 6, 7):
                cel.number_format = "[h]:mm:ss"
            if col in (8, 9):
                cel.number_format = "0.0%"
            if col == 10:
                cel.number_format = "#,##0.0"
            if linha % 2 == 0:
                cel.fill = fundo_faixa
        linha += 1
    linha_dados_fim = linha - 1
    if linha_dados_fim >= linha_dados_ini:
        re_.conditional_formatting.add(
            "E%d:E%d" % (linha_dados_ini, linha_dados_fim),
            DataBarRule(start_type="min", end_type="max", color=COR_ATRASO))
        re_.conditional_formatting.add(
            "H%d:H%d" % (linha_dados_ini, linha_dados_fim),
            DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                        color=COR_ATRASO))
    preparar_tabela(re_, linha_cab_rank, len(cab_rank), max(linha_dados_fim, linha_cab_rank))

    linha += 1
    linha = secao(re_, linha, "Piores 5 dias do período", 10)
    cab_piores = ["Data", "Dia da semana", "Placa", "Motorista", "1º horário", "Atraso"]
    cabecalho_tabela(re_, linha, cab_piores)
    linha += 1
    piores = sorted(dias, key=lambda d: -d["atraso_seg"])[:5]
    for d in piores:
        vals = [d["data_br"], d["dia_semana"], d["placa"],
                d["motorista"] if d["motorista"] else "Não identificado",
                _fracao_hora(d["primeira_saida_seg"]), _fracao_duracao(d["atraso_seg"])]
        for col, v in enumerate(vals, 1):
            cel = re_.cell(row=linha, column=col, value=v)
            cel.border = borda_cel
            cel.alignment = Alignment(horizontal="center" if col not in (2, 4) else "left",
                                      vertical="center", indent=1 if col == 4 else 0)
            if col == 5:
                cel.number_format = "hh:mm:ss"
            if col == 6:
                cel.number_format = "[h]:mm:ss"
                cel.font = Font(name=FONTE, bold=True, color=COR_ATRASO)
        linha += 1

    linha += 1
    rodape = [
        "Sem tolerância: 06:30 é 06:30 — a comparação é exata, ao segundo.",
        "Ignora saída após: %s — dias com a 1ª saída à noite não entram na conta (detalhes na aba Indicadores)." % (
            resumo["corte_noite"] if resumo["corte_noite"] else "desativado"),
        "Todas as abas de dados têm autofiltro pronto no cabeçalho — use pra recortar como quiser.",
    ]
    for txt in rodape:
        cel = re_.cell(row=linha, column=1, value="•  " + txt)
        cel.font = Font(name=FONTE, size=9, italic=True, color=COR_TEXTO_SUAVE)
        re_.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=10)
        linha += 1

    # ================================================== Saída de Veículos
    ws = wb.create_sheet("Saída de Veículos")
    cab = ["Data", "Dia da semana", "Placa do veículo", "Motorista",
           "Primeiro horário do dia", "Entrada do motorista",
           "Horas de atraso", "Situação", "Viagens", "KM do dia"]
    cabecalho_tabela(ws, 1, cab)
    for d in dias:
        ws.append([d["data_br"], d["dia_semana"], d["placa"],
                   d["motorista"] if d["motorista"] else "Não identificado",
                   _fracao_hora(d["primeira_saida_seg"]), _fracao_hora(d["entrada_seg"]),
                   _fracao_duracao(d["atraso_seg"]), d["situacao"], d["viagens"], d["km"]])
        linha = ws.max_row
        pinta = fundo_atraso if d["atraso_seg"] > 0 else fundo_ok
        for col in range(1, len(cab) + 1):
            cel = ws.cell(row=linha, column=col)
            cel.fill, cel.border = pinta, borda_cel
            cel.alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=linha, column=5).number_format = "hh:mm:ss"
        ws.cell(row=linha, column=6).number_format = "hh:mm:ss"
        ws.cell(row=linha, column=7).number_format = "[h]:mm:ss"
        ws.cell(row=linha, column=7).font = Font(
            name=FONTE, bold=True,
            color=COR_ATRASO if d["atraso_seg"] > 0 else COR_TEXTO_SUAVE)
        ws.cell(row=linha, column=4).alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.cell(row=linha, column=8).font = Font(
            name=FONTE, bold=True, color=COR_ATRASO if d["atraso_seg"] > 0 else COR_OK)
    larguras(ws, [12, 12, 16, 22, 15, 15, 14, 14, 9, 11])
    preparar_tabela(ws, 1, len(cab), ws.max_row)

    # =============================================== Consolidado por placa
    wp = wb.create_sheet("Consolidado por placa")
    cab2 = ["Placa", "Motorista", "Entrada", "Dias com saída", "Dias com atraso",
            "Horas de atraso (total)", "Atraso médio", "Maior atraso",
            "% saídas atrasadas", "Aderência %", "Viagens", "KM"]
    cabecalho_tabela(wp, 1, cab2)
    linha_ini = 2
    for p in resumo.get("por_placa", []):
        maior_placa = max((d["atraso_seg"] for d in dias if d["placa"] == p["placa"]), default=0)
        wp.append([p["placa"], p["motorista"] if p.get("motorista") else "Não identificado",
                   p["entrada"], p["dias"], p["dias_atrasados"],
                   _fracao_duracao(p["atraso_seg"]),
                   _fracao_duracao(p["atraso_seg"] / p["dias_atrasados"]) if p["dias_atrasados"] else 0,
                   _fracao_duracao(maior_placa), p["pct_atraso"] / 100.0, p["aderencia"] / 100.0,
                   p["viagens"], p["km"]])
        linha = wp.max_row
        for col in range(1, len(cab2) + 1):
            cel = wp.cell(row=linha, column=col)
            cel.border = borda_cel
            cel.alignment = Alignment(horizontal="center" if col not in (1, 2) else "left",
                                      vertical="center", indent=1 if col == 2 else 0)
            if linha % 2 == 0:
                cel.fill = fundo_faixa
        wp.cell(row=linha, column=1).font = Font(name=FONTE, bold=True, color=COR_MARCA)
        for col in (6, 7, 8):
            wp.cell(row=linha, column=col).number_format = "[h]:mm:ss"
        for col in (9, 10):
            wp.cell(row=linha, column=col).number_format = "0.0%"
        wp.cell(row=linha, column=9).font = Font(name=FONTE, bold=True, color=COR_ATRASO)
    linha_fim = wp.max_row
    if linha_fim >= linha_ini:
        wp.conditional_formatting.add("F%d:F%d" % (linha_ini, linha_fim),
                                      DataBarRule(start_type="min", end_type="max", color=COR_ATRASO))
        wp.conditional_formatting.add("I%d:I%d" % (linha_ini, linha_fim),
                                      DataBarRule(start_type="num", start_value=0,
                                                  end_type="num", end_value=1, color=COR_ATRASO))
    larguras(wp, [12, 22, 10, 14, 14, 18, 13, 13, 15, 12, 10, 11])
    preparar_tabela(wp, 1, len(cab2), max(linha_fim, 1))

    # ==================================================== Por dia da semana
    por_semana = montar_por_semana(dias)
    ws3 = wb.create_sheet("Por dia da semana")
    cab3 = ["Dia da semana", "Dias analisados", "Dias com atraso", "% atrasado",
            "Horas de atraso (total)", "Atraso médio", "Viagens", "KM"]
    cabecalho_tabela(ws3, 1, cab3)
    linha_ini3 = 2
    for p in por_semana:
        ws3.append([p["dia_semana"], p["dias"], p["dias_atrasados"], p["pct_atraso"] / 100.0,
                   _fracao_duracao(p["atraso_seg"]),
                   _fracao_duracao(p["atraso_seg"] / p["dias_atrasados"]) if p["dias_atrasados"] else 0,
                   p["viagens"], p["km"]])
        linha = ws3.max_row
        for col in range(1, len(cab3) + 1):
            cel = ws3.cell(row=linha, column=col)
            cel.border = borda_cel
            cel.alignment = Alignment(horizontal="left" if col == 1 else "center",
                                      vertical="center", indent=1 if col == 1 else 0)
            if linha % 2 == 0:
                cel.fill = fundo_faixa
        ws3.cell(row=linha, column=1).font = Font(name=FONTE, bold=True)
        ws3.cell(row=linha, column=4).number_format = "0.0%"
        for col in (5, 6):
            ws3.cell(row=linha, column=col).number_format = "[h]:mm:ss"
    linha_fim3 = ws3.max_row
    if linha_fim3 >= linha_ini3:
        ws3.conditional_formatting.add("D%d:D%d" % (linha_ini3, linha_fim3),
                                       DataBarRule(start_type="num", start_value=0,
                                                   end_type="num", end_value=1, color=COR_ATRASO))
    larguras(ws3, [15, 15, 15, 12, 18, 14, 10, 11])
    preparar_tabela(ws3, 1, len(cab3), max(linha_fim3, 1))

    # ========================================================= Indicadores
    wi = wb.create_sheet("Indicadores")
    wi.sheet_view.showGridLines = False
    larguras(wi, [36, 24])
    linha = secao(wi, 1, "Regras do indicador", 2)
    regras = [
        ("Entrada do motorista", resumo["entrada_padrao"]),
        ("Tolerância", "nenhuma — comparação exata ao segundo"),
        ("Ignora saída após", resumo["corte_noite"] if resumo["corte_noite"] else "desativado"),
    ]
    for nome, valor in regras:
        wi.append([nome, valor])
        wi.cell(row=wi.max_row, column=1).font = Font(name=FONTE, bold=True)
        wi.cell(row=wi.max_row, column=2).alignment = Alignment(horizontal="left")
    linha = wi.max_row + 2
    linha = secao(wi, linha, "Indicadores do período", 2)
    indicadores = [
        ("Período analisado", "%s a %s" % (resumo["periodo_inicio"], resumo["periodo_fim"])),
        ("Veículos", resumo["placas"]),
        ("Dias com saída", resumo["dias_analisados"]),
        ("Dias com atraso", resumo["dias_atrasados"]),
        ("Dias sem atraso", resumo["dias_sem_atraso"]),
        ("Horas de atraso (total)", resumo["atraso_total_hhmmss"]),
        ("Atraso médio (dias com atraso)", resumo["atraso_medio_hhmmss"]),
        ("Maior atraso do período", resumo["atraso_maior_hhmmss"]),
        ("Horário médio da 1ª saída", resumo["saida_media"]),
        ("Saída mais cedo", resumo["saida_mais_cedo"]),
        ("Saída mais tarde", resumo["saida_mais_tarde"]),
        ("Aderência ao horário (%)", resumo["aderencia"]),
        ("Total de viagens", resumo["viagens_total"]),
        ("KM total (trechos registrados)", resumo["km_total"]),
        ("Tempo em movimento", resumo["movimento_total_hhmmss"]),
        ("Tempo ocioso", resumo["ocioso_total_hhmmss"]),
    ]
    for nome, valor in indicadores:
        wi.append([nome, valor])
        wi.cell(row=wi.max_row, column=1).font = Font(name=FONTE, bold=True)

    # ---------------------------------------------------- impressão/PDF
    # paisagem + ajustada à largura, pra cada aba sair inteira numa folha
    # se alguém imprimir ou exportar pra PDF (senão o Excel quebra em
    # várias páginas por causa da quantidade de colunas)
    for aba in (re_, ws, wp, ws3, wi):
        aba.page_setup.orientation = "landscape"
        aba.page_setup.fitToWidth = 1
        aba.page_setup.fitToHeight = 0
        aba.sheet_properties.pageSetUpPr.fitToPage = True
        aba.page_margins.left = aba.page_margins.right = 0.35
        aba.page_margins.top = aba.page_margins.bottom = 0.5
        aba.page_margins.header = aba.page_margins.footer = 0.2

    wb.active = 0
    wb.save(destino)
    return destino


def main():
    arquivos = sorted(a for a in PASTA_PLANILHAS.glob("*.xls*")
                      if not a.name.startswith("~$"))
    if not arquivos:
        raise SystemExit("Nenhuma planilha encontrada em %s" % PASTA_PLANILHAS)

    print("Pasta: %s" % PASTA_PLANILHAS)
    registros = []
    for arq in arquivos:
        lidos = ler_arquivo(arq)
        placas = sorted(set(r["placa"] for r in lidos))
        print("  %s: %d deslocamentos  [%s]" % (arq.name, len(lidos),
                                                ", ".join(placas)))
        registros.extend(lidos)

    dias = montar_dias(registros)
    resumo = montar_resumo(dias)
    j = gravar_dados(dias, resumo, registros)
    x = gravar_excel(dias, resumo)

    print("")
    print("Periodo .................. %s a %s" % (resumo["periodo_inicio"], resumo["periodo_fim"]))
    print("Veiculos ................. %s" % resumo["placas"])
    print("Dias com saida ........... %s" % resumo["dias_analisados"])
    print("Dias com atraso .......... %s" % resumo["dias_atrasados"])
    print("Horas de atraso (total) .. %s" % resumo["atraso_total_hhmmss"])
    print("Atraso medio ............. %s" % resumo["atraso_medio_hhmmss"])
    print("Maior atraso ............. %s" % resumo["atraso_maior_hhmmss"])
    print("Horario medio da 1a saida  %s" % resumo["saida_media"])
    print("Aderencia ................ %s%%" % resumo["aderencia"])
    print("")
    print("Gerado: %s / %s" % (j.name, x.name))


if __name__ == "__main__":
    main()
